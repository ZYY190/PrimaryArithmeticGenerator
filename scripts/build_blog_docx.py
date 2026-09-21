from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
SOURCE = DOCS / "blog-draft.md"
OUTPUT = ROOT / "dist" / "博客园博文草稿.docx"
BODY_FONT = "Microsoft YaHei"
CODE_FONT = "Consolas"


def set_run_font(run, name=BODY_FONT, size=None, bold=None, color=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_shading(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def remove_paragraph_borders(style):
    p_pr = style.element.get_or_add_pPr()
    for child in list(p_pr):
        if child.tag == qn("w:pBdr"):
            p_pr.remove(child)


def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    run.append(r_pr)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, value):
    pattern = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|<https?://[^>]+>|\*[^*]+\*)")
    position = 0
    for match in pattern.finditer(value):
        if match.start() > position:
            run = paragraph.add_run(value[position:match.start()])
            set_run_font(run)
        token = match.group(0)
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, bold=True)
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name=CODE_FONT, size=9.5, color="A33A2B")
        elif token.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token).groups()
            add_hyperlink(paragraph, url, label)
        elif token.startswith("<") and token.endswith(">"):
            add_hyperlink(paragraph, token[1:-1], token[1:-1])
        elif token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, italic=True)
        position = match.end()
    if position < len(value):
        run = paragraph.add_run(value[position:])
        set_run_font(run)


def set_cell_text(cell, value, header=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    add_inline(paragraph, value)
    for run in paragraph.runs:
        set_run_font(run, size=9.5, bold=header or None)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if header:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "E8F0FF")
        tc_pr.append(shd)


def add_table(document, rows):
    if not rows:
        return
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            set_cell_text(table.cell(row_index, column_index), value.strip(), row_index == 0)
    document.add_paragraph()


def add_code_block(document, code_lines):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.35)
    paragraph.paragraph_format.right_indent = Cm(0.35)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    set_paragraph_shading(paragraph, "F3F5F7")
    run = paragraph.add_run("\n".join(code_lines))
    set_run_font(run, name=CODE_FONT, size=9, color="263238")
    return paragraph


def add_image(document, alt, path):
    picture_path = (DOCS / path).resolve()
    if not picture_path.exists():
        raise FileNotFoundError(picture_path)
    with Image.open(picture_path) as image:
        width, height = image.size
    max_width = Inches(6.2)
    if height / width > 1.35:
        max_width = Inches(5.6)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(picture_path), width=max_width)
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(8)
    caption_run = caption.add_run(alt)
    set_run_font(caption_run, size=9, color="6B7280", italic=True)


def parse_table(lines, start):
    rows = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        raw = lines[index].strip().strip("|")
        cells = [cell.strip() for cell in raw.split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells):
            rows.append(cells)
        index += 1
    return rows, index


def is_block_start(line):
    stripped = line.strip()
    return (
        not stripped
        or stripped.startswith("#")
        or stripped.startswith("```")
        or stripped.startswith("|")
        or stripped.startswith("![")
        or re.match(r"^[-*]\s+", stripped)
        or re.match(r"^\d+\.\s+", stripped)
    )


def build_docx():
    text = SOURCE.read_text(encoding="utf-8")
    lines = text.splitlines()
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.28
    normal.paragraph_format.space_after = Pt(6)
    for style_name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        style = styles[style_name]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        remove_paragraph_borders(style)
    styles["Title"].font.size = Pt(20)
    styles["Title"].font.bold = True
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 2"].font.size = Pt(13)
    styles["Heading 3"].font.size = Pt(11.5)
    styles["Heading 1"].font.color.rgb = RGBColor(31, 56, 100)
    styles["Heading 2"].font.color.rgb = RGBColor(45, 78, 134)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("小学四则运算自动出题程序  |  第 ")
    set_run_font(footer_run, size=8.5, color="7A8494")
    page_field = OxmlElement("w:fldSimple")
    page_field.set(qn("w:instr"), "PAGE")
    footer._p.append(page_field)
    end_run = footer.add_run(" 页")
    set_run_font(end_run, size=8.5, color="7A8494")

    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            index += 1
            continue
        if stripped.startswith("```"):
            index += 1
            code_lines = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            add_code_block(document, code_lines)
            index += 1
            continue
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            value = stripped[level:].strip()
            if level == 1:
                paragraph = document.add_paragraph(style="Title")
            else:
                paragraph = document.add_paragraph(style=f"Heading {min(level - 1, 3)}")
            add_inline(paragraph, value)
            index += 1
            continue
        if stripped.startswith("|"):
            rows, index = parse_table(lines, index)
            add_table(document, rows)
            continue
        image_match = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if image_match:
            add_image(document, image_match.group(1), image_match.group(2))
            index += 1
            continue
        bullet_match = re.match(r"^[-*]\s+(.*)$", stripped)
        number_match = re.match(r"^\d+\.\s+(.*)$", stripped)
        if bullet_match or number_match:
            match = bullet_match or number_match
            paragraph = document.add_paragraph(style="List Bullet" if bullet_match else "List Number")
            add_inline(paragraph, match.group(1))
            index += 1
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines) and not is_block_start(lines[index]):
            paragraph_lines.append(lines[index].strip())
            index += 1
        paragraph = document.add_paragraph()
        if len(paragraph_lines) == 1:
            add_inline(paragraph, paragraph_lines[0])
        else:
            for line_index, value in enumerate(paragraph_lines):
                if line_index:
                    paragraph.add_run().add_break()
                add_inline(paragraph, value)

    document.core_properties.title = "第一次个人编程作业：小学四则运算题目生成器"
    document.core_properties.author = "ZYY190"
    document.core_properties.subject = "软件工程个人编程作业博客草稿"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_docx()
